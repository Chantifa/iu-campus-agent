import json

import pytest

from iu_agent.rag.loaders import EmptyDocumentError, UnsupportedFileError, html_to_text, load_file


def test_pdf_pages(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Hello Advanced Maths page one")
    doc.new_page().insert_text((72, 72), "Second page about eigenvalues")
    target = tmp_path / "a.pdf"
    doc.save(str(target))
    doc.close()

    sections = load_file(target)
    assert len(sections) == 2
    assert sections[0].meta == {"page": 1, "pages": 2}
    assert "eigenvalues" in sections[1].text


def test_docx_headings_and_tables(tmp_path):
    from docx import Document

    document = Document()
    document.add_heading("Course Title", level=1)
    document.add_paragraph("Body text about gradient descent.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "left"
    table.rows[0].cells[1].text = "right"
    target = tmp_path / "a.docx"
    document.save(str(target))

    [section] = load_file(target)
    assert "# Course Title" in section.text
    assert "gradient descent" in section.text
    assert "left | right" in section.text


def test_pptx_slides(tmp_path):
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Slide title"
    slide.placeholders[1].text = "A bullet point"
    target = tmp_path / "a.pptx"
    presentation.save(str(target))

    [section] = load_file(target)
    assert section.meta["slide"] == 1
    assert "Slide title" in section.text and "A bullet point" in section.text


def test_ipynb_cells(tmp_path):
    notebook = {
        "cells": [
            {"cell_type": "markdown", "source": ["# Exercise 1\n", "Explain backpropagation."]},
            {
                "cell_type": "code",
                "source": ["print('hi')"],
                "outputs": [{"output_type": "stream", "text": ["hi\n"]}],
            },
        ]
    }
    target = tmp_path / "a.ipynb"
    target.write_text(json.dumps(notebook), encoding="utf-8")

    [section] = load_file(target)
    assert "# Exercise 1" in section.text
    assert "```python" in section.text
    assert "Output:\nhi" in section.text


def test_html_strips_scripts(tmp_path):
    target = tmp_path / "a.html"
    target.write_text(
        "<html><body><h1>Head</h1><script>alert(1)</script><p>Para</p></body></html>", encoding="utf-8"
    )
    [section] = load_file(target)
    assert "# Head" in section.text
    assert "Para" in section.text
    assert "alert" not in section.text


def test_html_to_text_lists():
    text = html_to_text("<ul><li>one</li><li>two</li></ul>")
    assert "- one" in text and "- two" in text


def test_unsupported_and_empty(tmp_path):
    binary = tmp_path / "a.exe"
    binary.write_bytes(b"\x00\x01")
    with pytest.raises(UnsupportedFileError):
        load_file(binary)
    empty = tmp_path / "empty.txt"
    empty.write_text("   \n", encoding="utf-8")
    with pytest.raises(EmptyDocumentError):
        load_file(empty)
