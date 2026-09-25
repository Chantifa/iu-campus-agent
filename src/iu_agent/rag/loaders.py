"""Turn course files into text sections (one per page, slide, sheet or whole document).

Supported: PDF, DOCX, PPTX, XLSX, HTML, Jupyter notebooks and plain text formats
(Markdown, TXT, LaTeX, BibTeX, Python, CSV, reStructuredText).
"""

from __future__ import annotations

import json
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

SUPPORTED_EXTENSIONS: frozenset[str] = frozenset(
    {
        ".pdf",
        ".docx",
        ".pptx",
        ".xlsx",
        ".md",
        ".markdown",
        ".rst",
        ".txt",
        ".html",
        ".htm",
        ".ipynb",
        ".tex",
        ".bib",
        ".py",
        ".csv",
    }
)

_TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "cp1252", "latin-1")
_WS_RE = re.compile(r"[ \t ]+")
_BLANK_RE = re.compile(r"\n{3,}")


class UnsupportedFileError(ValueError):
    pass


class EmptyDocumentError(ValueError):
    pass


@dataclass
class Section:
    """A piece of extracted text plus locator metadata (page, slide, sheet ...)."""

    text: str
    meta: dict[str, Any] = field(default_factory=dict)


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = _WS_RE.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _BLANK_RE.sub("\n\n", text).strip()


def read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in _TEXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def html_to_text(html: str) -> str:
    """Convert HTML (Moodle pages, exported course pages) to Markdown-ish plain text."""
    from bs4 import BeautifulSoup
    from markdownify import markdownify

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "iframe", "svg"]):
        tag.decompose()
    return clean_text(markdownify(str(soup), heading_style="ATX", bullets="-"))


# ----------------------------------------------------------------------------- loaders
def load_pdf(path: Path) -> list[Section]:
    try:
        import pymupdf
    except ImportError:  # pragma: no cover - older package name
        import fitz as pymupdf  # type: ignore[no-redef]

    sections: list[Section] = []
    with pymupdf.open(str(path)) as doc:
        total = doc.page_count
        for index, page in enumerate(doc, start=1):
            text = clean_text(page.get_text("text"))
            if text:
                sections.append(Section(text, {"page": index, "pages": total}))
    return sections


def load_docx(path: Path) -> list[Section]:
    from docx import Document as DocxDocument

    document = DocxDocument(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        if style.startswith("heading"):
            digits = "".join(ch for ch in style if ch.isdigit())
            level = min(int(digits), 6) if digits else 2
            parts.append("#" * level + " " + text)
        elif style.startswith("title"):
            parts.append("# " + text)
        else:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = clean_text("\n".join(parts))
    return [Section(text)] if text else []


def load_pptx(path: Path) -> list[Section]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    slides = list(presentation.slides)
    sections: list[Section] = []
    for index, slide in enumerate(slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        lines.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append("Notes: " + notes)
        text = clean_text("\n".join(lines))
        if text:
            sections.append(Section(text, {"slide": index, "slides": len(slides)}))
    return sections


def load_xlsx(path: Path, max_rows: int = 2000) -> list[Section]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    sections: list[Section] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= max_rows:
                    rows.append(f"... ({sheet.max_row - max_rows} more rows)")
                    break
                cells = ["" if value is None else str(value) for value in row]
                if any(cell.strip() for cell in cells):
                    rows.append(" | ".join(cells).rstrip(" |"))
            text = clean_text("\n".join(rows))
            if text:
                sections.append(Section(text, {"sheet": sheet.title}))
    finally:
        workbook.close()
    return sections


def load_html(path: Path) -> list[Section]:
    text = html_to_text(read_text(path))
    return [Section(text)] if text else []


def load_ipynb(path: Path) -> list[Section]:
    notebook = json.loads(read_text(path))
    parts: list[str] = []
    for index, cell in enumerate(notebook.get("cells", []), start=1):
        source = cell.get("source", "")
        if isinstance(source, list):
            source = "".join(source)
        source = source.strip()
        if not source:
            continue
        cell_type = cell.get("cell_type")
        if cell_type == "markdown":
            parts.append(source)
        elif cell_type == "code":
            block = f"```python\n{source}\n```"
            outputs: list[str] = []
            for output in cell.get("outputs", []) or []:
                if output.get("output_type") == "stream":
                    outputs.append("".join(output.get("text", [])))
                elif "data" in output and "text/plain" in output["data"]:
                    data = output["data"]["text/plain"]
                    outputs.append("".join(data) if isinstance(data, list) else str(data))
            if outputs:
                joined = "\n".join(outputs).strip()
                block += "\nOutput:\n" + (joined[:1500] + " ..." if len(joined) > 1500 else joined)
            parts.append(f"[cell {index}]\n{block}")
    text = clean_text("\n\n".join(parts))
    return [Section(text)] if text else []


def load_text(path: Path) -> list[Section]:
    text = clean_text(read_text(path))
    return [Section(text)] if text else []


LOADERS: dict[str, Callable[[Path], list[Section]]] = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".pptx": load_pptx,
    ".xlsx": load_xlsx,
    ".html": load_html,
    ".htm": load_html,
    ".ipynb": load_ipynb,
}


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_EXTENSIONS


def load_file(path: Path) -> list[Section]:
    """Extract text sections from ``path``; raises for unsupported or empty (e.g. scanned) files."""
    path = Path(path)
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(f"Unsupported file type: {path.name}")
    loader = LOADERS.get(extension, load_text)
    sections = loader(path)
    if not sections:
        raise EmptyDocumentError(f"No extractable text in {path.name} (scanned document or empty file)")
    return sections
